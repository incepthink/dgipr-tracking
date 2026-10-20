"""Roll-ups over a form's numeric fields (aggregate=sum), and the Excel export.

Every function takes the office codes the caller may see (app.forms.service.scope) and sums in SQL. Only
submitted / accepted / returned periods count (service.COUNTED); drafts are work in progress. Totals sum
entered values only; how complete they are is reported separately (coverage, compliance), so a missing
period is never silently read as zero.

The five sheet tables reproduce the admin's workbook (दररोजचा अहवाल, विभागनिहाय एकत्रित, तारखेनिहाय अहवाल,
जिल्हानिहाय दररोजचा अहवाल, जिल्हानिहाय एकत्रित); workbook() writes them after सूचना and डेटा एंट्री, with
admin-corrected cells highlighted (the workbook's own yellow) and explained in a cell comment."""
from __future__ import annotations

import csv
import io
from collections import defaultdict
from datetime import date, timedelta
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.forms.periods import label as period_label_of, period_containing, period_end, periods
from app.forms.schema import active_fields, sum_fields
from app.forms.sheet_names import division_name, office_names
from app.forms.service import COUNTED, _office_meta, _original_of, _subs, value_of
from app.models import Division, Form, FormSubmission, FormValue, Office
from app.reporting.marathi import long_date, period_label

HIGHLIGHT = "FFF2CC"  # the fill the admin's workbook uses for corrected cells
HEADER_FILL = "1F4E78"


def _cols(form: Form) -> list:
    return sum_fields(form)


def _zero(form: Form) -> dict[str, float]:
    return {f.key: 0 for f in _cols(form)}


def _num(v):
    return int(v) if v is not None and float(v).is_integer() else v


def _where(form: Form, offices: list[str], start: date, end: date):
    return (FormSubmission.form_id == form.id, FormSubmission.office_code.in_(offices or ["-"]),
            FormSubmission.period_start >= start, FormSubmission.period_start <= end,
            FormSubmission.status.in_(COUNTED), FormValue.value_num.isnot(None),
            FormValue.field_id.in_([f.id for f in _cols(form)] or [-1]))


def _q(form: Form, *cols):
    return select(*cols).select_from(FormValue).join(FormSubmission, FormSubmission.id == FormValue.submission_id)


def _keys(form: Form) -> dict[int, str]:
    return {f.id: f.key for f in _cols(form)}


def totals(s: Session, form: Form, offices: list[str], start: date, end: date) -> dict[str, float]:
    out = _zero(form)
    keys = _keys(form)
    for fid, total in s.execute(_q(form, FormValue.field_id, func.sum(FormValue.value_num))
                                .where(*_where(form, offices, start, end)).group_by(FormValue.field_id)).all():
        out[keys[fid]] = _num(total or 0)
    return out


def coverage(s: Session, form: Form, offices: list[str], start: date, end: date) -> dict[str, Any]:
    n_periods = len(periods(form, start, end))
    filled = s.scalar(_q(form, func.count(FormValue.id)).where(*_where(form, offices, start, end))) or 0
    reporting = s.scalar(_q(form, func.count(func.distinct(FormSubmission.office_code)))
                         .where(*_where(form, offices, start, end))) or 0
    expected = len(offices) * n_periods * len(_cols(form))
    return {"filled_cells": int(filled), "expected_cells": expected,
            "pct": round(100 * filled / expected, 1) if expected else 0.0,
            "offices": len(offices), "offices_reporting": int(reporting), "days": n_periods}


def summary(s: Session, form: Form, offices: list[str], start: date, end: date) -> dict[str, Any]:
    """Totals for the period and for the previous period of equal length."""
    n = (end - start).days + 1
    p_end = start - timedelta(days=1)
    p_start = p_end - timedelta(days=n - 1)
    return {"start": start.isoformat(), "end": end.isoformat(),
            "previous": {"start": p_start.isoformat(), "end": p_end.isoformat(),
                         "totals": totals(s, form, offices, p_start, p_end)},
            "totals": totals(s, form, offices, start, end),
            "coverage": coverage(s, form, offices, start, end)}


def breakdown(s: Session, form: Form, offices: list[str], start: date, end: date,
              level: str = "office") -> dict[str, Any]:
    """Per office (जिल्हानिहाय) or per division (विभागनिहाय) sums, plus the total row."""
    if level not in ("office", "division"):
        raise ValueError("level must be office or division")
    meta = _office_meta(s, offices)
    keys = _keys(form)
    group_col = Office.division_code if level == "division" else FormSubmission.office_code
    q = (_q(form, group_col, FormValue.field_id, func.sum(FormValue.value_num), func.count(FormValue.id))
         .join(Office, Office.code == FormSubmission.office_code)
         .where(*_where(form, offices, start, end)).group_by(group_col, FormValue.field_id))
    sums: dict[str, dict[str, float]] = defaultdict(lambda: _zero(form))
    cells: dict[str, int] = defaultdict(int)
    entered: dict[str, dict[str, int]] = defaultdict(lambda: {k: 0 for k in keys.values()})
    for code, fid, total, n in s.execute(q).all():
        sums[code][keys[fid]] = _num(total or 0)
        cells[code] += int(n)
        entered[code][keys[fid]] = int(n)
    per = len(periods(form, start, end)) * len(keys)
    rows: list[dict[str, Any]] = []
    if level == "office":
        for o, d in meta:
            rows.append({"code": o.code, "name_mr": o.name_mr, "name_en": o.name_en, "kind": o.kind,
                         "division": d.code, "division_mr": d.summary_label_mr, "values": sums[o.code],
                         "entered": entered[o.code], "filled_cells": cells[o.code], "expected_cells": per})
    else:
        per_div: dict[str, list[Office]] = defaultdict(list)
        divs: dict[str, Division] = {}
        for o, d in meta:
            per_div[d.code].append(o)
            divs[d.code] = d
        for code, offs in per_div.items():
            rows.append({"code": code, "name_mr": divs[code].summary_label_mr,
                         "name_en": code.replace("_", " ").title(), "values": sums[code], "entered": entered[code],
                         "filled_cells": cells[code], "expected_cells": per * len(offs), "offices": len(offs)})
    total = _zero(form)
    for r in rows:
        for k in total:
            total[k] += r["values"][k]
    return {"level": level, "start": start.isoformat(), "end": end.isoformat(), "rows": rows, "total": total}


def series(s: Session, form: Form, offices: list[str], start: date, end: date) -> list[dict[str, Any]]:
    """Per-period (तारखेनिहाय) totals; periods with no entries are 0 with entered_cells=0, so a chart can
    show them as gaps. `entered` counts the values behind each field's total (0 = nothing entered)."""
    keys = _keys(form)
    by: dict[date, dict[str, float]] = defaultdict(lambda: _zero(form))
    entered: dict[date, int] = defaultdict(int)
    per_field: dict[date, dict[str, int]] = defaultdict(lambda: {k: 0 for k in keys.values()})
    for ps, fid, total, n in s.execute(
            _q(form, FormSubmission.period_start, FormValue.field_id, func.sum(FormValue.value_num),
               func.count(FormValue.id))
            .where(*_where(form, offices, start, end))
            .group_by(FormSubmission.period_start, FormValue.field_id)).all():
        by[ps][keys[fid]] = _num(total or 0)
        entered[ps] += int(n)
        per_field[ps][keys[fid]] = int(n)
    return [{"date": ps.isoformat(), "label": period_label_of(form, ps), "values": by[ps],
             "entered": per_field[ps], "entered_cells": entered[ps], "expected_cells": len(offices) * len(keys)}
            for ps in periods(form, start, end)]


def rankings(s: Session, form: Form, offices: list[str], start: date, end: date, column: str,
             level: str = "office") -> dict[str, Any]:
    if column not in _zero(form):
        raise ValueError(f"unknown column {column}")
    b = breakdown(s, form, offices, start, end, level)
    total = b["total"][column]
    rows = sorted(({"code": r["code"], "name_mr": r["name_mr"], "name_en": r["name_en"],
                    "division_mr": r.get("division_mr"), "value": r["values"][column],
                    "share": round(100 * r["values"][column] / total, 1) if total else 0.0} for r in b["rows"]),
                  key=lambda r: (-r["value"], r["name_en"]))
    return {"column": column, "level": level, "total": total, "rows": rows}


# --------------------------------------------------------------------------- sheet-equivalent tables
TABLES = {
    "daily_division": "दररोजचा अहवाल (विभागनिहाय)",
    "daily_office": "जिल्हानिहाय दररोजचा अहवाल",
    "cumulative_division": "विभागनिहाय एकत्रित",
    "cumulative_office": "जिल्हानिहाय एकत्रित",
    "datewise": "तारखेनिहाय अहवाल",
}
# Sheet names and order in the admin's workbook.
SHEETS = (("daily_division", "दररोजचा अहवाल"), ("cumulative_division", "विभागनिहाय एकत्रित"),
          ("datewise", "तारखेनिहाय अहवाल"), ("daily_office", "जिल्हानिहाय दररोजचा अहवाल"),
          ("cumulative_office", "जिल्हानिहाय एकत्रित"))


def table(s: Session, form: Form, offices: list[str], start: date, end: date, kind: str) -> dict[str, Any]:
    """{title, subtitle, header, rows, total, text_cols, blank} for one of TABLES, laid out like the sheet. The
    daily tables show the single period containing `start`. blank[i][j] is true when nothing at all was entered
    behind value cell (i, j) — the sheet shows 0 there, the screen shows it as not reported."""
    if kind not in TABLES:
        raise ValueError(f"unknown table {kind}")
    if kind.startswith("daily_"):
        start = period_containing(form, start) or start
        end = start
    cols = _cols(form)
    labels = [f.label_mr for f in cols]
    keys = [f.key for f in cols]
    if kind == "datewise":
        header = ["दिनांक", *labels]
        ser = series(s, form, offices, start, end)
        rows = [[r["label"], *[r["values"][k] for k in keys]] for r in ser]
        blank = [[not r["entered"][k] for k in keys] for r in ser]
        total = ["एकूण", *[_num(sum(r[i + 1] for r in rows)) for i in range(len(keys))]]
    else:
        level = "division" if kind.endswith("division") else "office"
        b = breakdown(s, form, offices, start, end, level)
        blank = [[not r["entered"][k] for k in keys] for r in b["rows"]]
        if level == "division":
            header = ["विभागाचे नाव", *labels]
            members: dict[str, list[str]] = defaultdict(list)
            for o, _ in _office_meta(s, offices):
                members[o.division_code].append(o.code)
            rows = [[division_name(r["code"], members[r["code"]], r["name_mr"]), *[r["values"][k] for k in keys]]
                    for r in b["rows"]]
            total = ["राज्य एकूण", *[b["total"][k] for k in keys]]
        else:
            header = ["विभागाचे नाव", "जिल्ह्याचे नाव", *labels]
            rows = [[*office_names(r["code"], r["division_mr"], r["name_mr"]), *[r["values"][k] for k in keys]]
                    for r in b["rows"]]
            total = ["", "एकत्रित एकूण", *[b["total"][k] for k in keys]]
    shown_end = period_end(form, start) if kind.startswith("daily_") else end
    title = f"{form.title_mr} — {TABLES[kind]}".strip(" —")
    return {"kind": kind, "title": title, "subtitle": f"दि. {period_label(start, shown_end)}",
            "start": start.isoformat(), "end": end.isoformat(), "header": header, "rows": rows, "total": total,
            "text_cols": len(header) - len(keys), "blank": blank}


def to_csv(t: dict[str, Any]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([t["title"]])
    w.writerow([t["subtitle"]])
    w.writerow(t["header"])
    w.writerows(t["rows"])
    w.writerow(t["total"])
    return "﻿" + buf.getvalue()  # BOM so Excel opens Devanagari correctly


# --------------------------------------------------------------------------- xlsx
def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color="999999")
    return {"box": Border(left=thin, right=thin, top=thin, bottom=thin),
            "head_font": Font(bold=True, color="FFFFFF"), "head_fill": PatternFill("solid", fgColor=HEADER_FILL),
            "title_font": Font(bold=True, size=13, color="FFFFFF"), "bold": Font(bold=True),
            "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
            "wrap": Alignment(horizontal="left", vertical="top", wrap_text=True),
            "highlight": PatternFill("solid", fgColor=HIGHLIGHT), "total_fill": PatternFill("solid", fgColor="E8ECF4")}


def _write_table(ws, t: dict[str, Any], *, header_row: int = 3) -> None:
    """Title in row 1, subtitle in row 2, header at header_row, then rows and the bold total row."""
    from openpyxl.utils import get_column_letter
    st = _styles()
    width = len(t["header"])
    ws.cell(1, 1, t["title"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.cell(1, 1).font, ws.cell(1, 1).fill, ws.cell(1, 1).alignment = st["title_font"], st["head_fill"], st["center"]
    ws.cell(2, 1, t["subtitle"]).font = st["bold"]
    for i, h in enumerate(t["header"], start=1):
        c = ws.cell(header_row, i, h)
        c.font, c.fill, c.border, c.alignment = st["head_font"], st["head_fill"], st["box"], st["center"]
    r = header_row
    for row in t["rows"]:
        r += 1
        for i, v in enumerate(row, start=1):
            ws.cell(r, i, v).border = st["box"]
    r += 1
    for i, v in enumerate(t["total"], start=1):
        c = ws.cell(r, i, v)
        c.font, c.border, c.fill = st["bold"], st["box"], st["total_fill"]
    tc = t["text_cols"]
    for i in range(1, width + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22 if i <= tc else 12
    ws.row_dimensions[header_row].height = 48
    ws.freeze_panes = ws.cell(header_row + 1, tc + 1)


def to_xlsx(t: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = TABLES[t["kind"]][:31]
    _write_table(ws, t)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_out(f, v):
    if v is None:
        return None
    if f.type == "yesno":
        return {"yes": "होय", "no": "नाही"}.get(v, v)
    if f.type == "date":
        try:
            return date.fromisoformat(v)
        except ValueError:
            return v
    return v


def workbook(s: Session, form: Form, offices: list[str], start: date, end: date) -> bytes:
    """The admin's workbook: सूचना, डेटा एंट्री (every office × period, corrected cells highlighted with a
    comment), then the five roll-up sheets (the daily ones for the last period of the range)."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.utils import get_column_letter

    st = _styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "सूचना"
    ws.column_dimensions["A"].width = 120
    ws.cell(1, 1, f"{form.title_mr} — वापर सूचना").font = st["title_font"]
    ws.cell(1, 1).fill = st["head_fill"]
    r = 3
    for line in (form.instructions_mr or "").splitlines() or [""]:
        ws.cell(r, 1, line).alignment = st["wrap"]
        r += 1
    r += 1
    ws.cell(r, 1, f"कालावधी: {period_label(start, end)} · निर्यात: केवळ सादर केलेली माहिती; पिवळ्या रंगातील "
                  "आकडे राज्य कार्यालयाने दुरुस्त केलेले आहेत (मूळ आकडा टिप्पणीत).").font = st["bold"]

    # ------------------------------------------------------------- डेटा एंट्री
    de = wb.create_sheet("डेटा एंट्री")
    fields = active_fields(form)
    ps_list = periods(form, start, end)
    subs = _subs(s, form, offices, ps_list)
    meta = _office_meta(s, offices)
    header = ["दिनांक", "विभागाचे नाव", "जिल्ह्याचे नाव", *[f.label_mr for f in fields], "स्थिती"]
    width = len(header)
    de.cell(1, 1, f"{form.title_mr} — एकत्रित डेटा एंट्री")
    de.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    de.cell(1, 1).font, de.cell(1, 1).fill, de.cell(1, 1).alignment = st["title_font"], st["head_fill"], st["center"]
    de.cell(2, 1, f"दि. {period_label(start, end)}").font = st["bold"]
    for i, h in enumerate(header, start=1):
        c = de.cell(3, i, h)
        c.font, c.fill, c.border, c.alignment = st["head_font"], st["head_fill"], st["box"], st["center"]
    status_mr = {"submitted": "सादर", "accepted": "स्वीकृत", "returned": "परत पाठवले"}
    row = 3
    for ps in ps_list:
        for o, d in meta:
            row += 1
            sub = subs.get((o.code, ps))
            counted = sub is not None and sub.status in COUNTED
            de.cell(row, 1, ps).number_format = "yyyy-mm-dd"
            div_name, office_name = office_names(o.code, d.summary_label_mr, o.name_mr)
            de.cell(row, 2, div_name)
            de.cell(row, 3, office_name)
            vals = {v.field_id: v for v in sub.values} if counted else {}
            for i, f in enumerate(fields, start=4):
                v = vals.get(f.id)
                c = de.cell(row, i, _cell_out(f, value_of(f, v)) if v is not None else None)
                if f.type == "date" and isinstance(c.value, date):
                    c.number_format = "yyyy-mm-dd"
                if v is not None and v.corrected_at is not None:
                    orig = _original_of(f, v)
                    orig = "रिकामे" if orig is None else _cell_out(f, orig)
                    c.fill = st["highlight"]
                    c.comment = Comment(f"मूळ आकडा: {orig}\nदुरुस्ती: {v.corrected_by}, "
                                        f"{long_date(v.corrected_at.date())}\nकारण: {v.correction_note}",
                                        v.corrected_by or "admin", width=260, height=110)
            de.cell(row, width, status_mr.get(sub.status, "") if counted else ("मसुदा" if sub else ""))
    for i in range(1, width + 1):
        de.column_dimensions[get_column_letter(i)].width = 14 if i == 1 else 20 if i <= 3 else 12
    de.row_dimensions[3].height = 48
    de.freeze_panes = "D4"

    # ------------------------------------------------------------- roll-ups
    last = ps_list[-1] if ps_list else start
    for kind, name in SHEETS:
        t = table(s, form, offices, last if kind.startswith("daily_") else start,
                  last if kind.startswith("daily_") else end, kind)
        _write_table(wb.create_sheet(name), t, header_row=4 if kind.startswith("daily_") else 3)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
