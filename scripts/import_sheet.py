"""Import the admin's data-entry workbook ("दुष्काळ_प्रसिद्धी_अहवाल_….xlsx", sheet "डेटा एंट्री") into a form.

Header on row 3, data from row 4: दिनांक, विभागाचे नाव, जिल्ह्याचे नाव, then the 12 report columns in COLUMNS
order (the form's fields with those keys). Only filled cells are loaded: a blank cell stays missing, a typed 0
is a confirmed zero. Each office-day becomes a submitted period; existing values for the same office-day-field
are overwritten (the import is re-runnable). One audit entry records the batch, with actor "import:xlsx".

  python scripts/import_sheet.py PATH.xlsx --dry-run      # totals only, no writes
  python scripts/import_sheet.py PATH.xlsx [--form drought_measures]
"""
from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

ACTOR = "import:xlsx"
SHEET = "डेटा एंट्री"
HEADER_ROW, FIRST_ROW = 3, 4

from app.forms.sheet_names import OFFICE_MAP  # noqa: E402  (sheet spellings -> office code)


def read_sheet(path: Path) -> tuple[list[tuple[date, str, list[int | None]]], list[str]]:
    """[(day, office_code, 12 values or None)] for rows with at least one value, plus problems found."""
    from openpyxl import load_workbook

    from app.reporting.columns import COLUMN_KEYS

    ws = load_workbook(path, read_only=True, data_only=True)[SHEET]
    out, problems = [], []
    for n, row in enumerate(ws.iter_rows(min_row=FIRST_ROW, values_only=True), start=FIRST_ROW):
        if not row or row[0] is None:
            continue
        values = list(row[3:3 + len(COLUMN_KEYS)])
        values += [None] * (len(COLUMN_KEYS) - len(values))
        if all(v is None for v in values):
            continue
        day = row[0].date() if isinstance(row[0], datetime) else row[0]
        division, office = (str(row[1] or "").strip(), str(row[2] or "").strip())
        code = OFFICE_MAP.get((division, office))
        if not isinstance(day, date):
            problems.append(f"row {n}: date {row[0]!r} is not a date")
            continue
        if code is None:
            problems.append(f"row {n}: unknown office {division} / {office}")
            continue
        clean: list[int | None] = []
        for k, v in zip(COLUMN_KEYS, values):
            if v is None or (isinstance(v, str) and not v.strip()):
                clean.append(None)
            elif isinstance(v, (int, float)) and float(v).is_integer() and v >= 0:
                clean.append(int(v))
            else:
                problems.append(f"row {n}: {office} {day} {k}: {v!r} is not a whole number ≥ 0")
                clean.append(None)
        out.append((day, code, clean))
    return out, problems


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--form", default="drought_measures", help="form slug")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")

    from sqlalchemy import select

    from app import audit
    from app.db import init_db, make_engine, make_session_factory
    from app.forms.service import save
    from app.models import Form
    from app.registry.loader import bootstrap
    from app.reporting.columns import COLUMN_KEYS
    from app.settings import load_settings

    rows, problems = read_sheet(args.xlsx)
    for p in problems:
        print("PROBLEM:", p)
    totals = [sum((vals[i] or 0) for _, _, vals in rows) for i in range(len(COLUMN_KEYS))]
    per_office: dict[str, list[int]] = defaultdict(lambda: [0] * len(COLUMN_KEYS))
    for _, code, vals in rows:
        for i, v in enumerate(vals):
            per_office[code][i] += v or 0
    days = sorted({d for d, _, _ in rows})
    print(f"{len(rows)} office-day rows with values, {len(days)} days "
          f"({days[0] if days else '-'} … {days[-1] if days else '-'}), {len(per_office)} offices")
    print("state total (राज्य एकूण):", totals)
    if problems:
        print(f"{len(problems)} problem(s) — fix the workbook or OFFICE_MAP; nothing imported.")
        sys.exit(1)
    if args.dry_run:
        return

    settings = load_settings()
    engine = make_engine(settings.database_url)
    init_db(engine)
    factory = make_session_factory(engine)
    with factory() as s:
        bootstrap(s, settings.config_dir)
        form = s.scalar(select(Form).where(Form.slug == args.form))
        if form is None:
            sys.exit(f"form {args.form!r} not found")
        missing = [k for k in COLUMN_KEYS if k not in {f.key for f in form.fields}]
        if missing:
            sys.exit(f"form {form.slug} has no field(s) {', '.join(missing)}")
        for day, code, vals in rows:
            values = {k: v for k, v in zip(COLUMN_KEYS, vals) if v is not None}
            save(s, None, form, code, day, values, confirm_zero=True, actor=ACTOR, check=False, submit=True,
                 note=f"imported from {args.xlsx.name}")
        audit.record(s, ACTOR, "form.import", "form", form.id,
                     after={"file": args.xlsx.name, "rows": len(rows),
                            "days": [days[0].isoformat(), days[-1].isoformat()] if days else None,
                            "totals": dict(zip(COLUMN_KEYS, totals))})
        s.commit()
        print(f"imported into form {form.slug} (id {form.id}): {len(rows)} office-day(s)")


if __name__ == "__main__":
    main()
