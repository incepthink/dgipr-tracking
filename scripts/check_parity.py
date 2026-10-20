"""Excel parity: our five sheet tables (app.forms.analytics.table) against the admin's own workbook.

  python scripts/check_parity.py "PATH/दुष्काळ_….xlsx" [--start 2026-10-09 --end 2026-10-15] [--no-libreoffice]

1. The workbook's डेटा एंट्री is imported into a scratch database (a temp directory; the real one is untouched),
   the same way scripts/import_sheet.py does it.
2. The workbook's own five views are worked out twice where possible:
     - LibreOffice headless recalculates a copy (per day for the daily sheets, which read their date from B2), and
     - an independent SUMIFS over डेटा एंट्री that copies each view's formula criteria (exact names, as Excel does).
   The two must agree; LibreOffice is skipped when it isn't installed (or with --no-libreoffice).
3. Every value cell is compared with table(): the daily views for each day of [start, end], the date-wise view's
   rows for those days, and the cumulative views over the workbook's whole range (which is what its SUMIFS sum).
   A difference is "explained" when a डेटा एंट्री name differs from the view's label (so the workbook's SUMIFS drops
   it) but maps to the same office (sheet_names.OFFICE_MAP) — the name problem the yellow corrections fixed.
   Anything else is unexplained and the script exits 1."""
from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

DE = "डेटा एंट्री"
# view -> (sheet, header row, text columns, the डेटा एंट्री columns its SUMIFS matches on)
VIEWS = {
    "daily_division": ("दररोजचा अहवाल", 4, 1, ("date", "division")),
    "cumulative_division": ("विभागनिहाय एकत्रित", 3, 1, ("division",)),
    "datewise": ("तारखेनिहाय अहवाल", 3, 1, ("date",)),
    "daily_office": ("जिल्हानिहाय दररोजचा अहवाल", 4, 2, ("date", "division", "office")),
    "cumulative_office": ("जिल्हानिहाय एकत्रित", 3, 2, ("division", "office")),
}
TOTAL_LABELS = {"एकूण", "राज्य एकूण", "एकत्रित एकूण", "एकुण"}


def _day(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(str(v)[:10])
    except ValueError:
        return None


def _s(v) -> str:
    return "" if v is None else str(v)


# --------------------------------------------------------------------------- the workbook's side
def data_rows(path: Path) -> list[tuple[date, str, str, list]]:
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True)[DE]
    out = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or row[0] is None:
            continue
        d = _day(row[0])
        if d:
            out.append((d, _s(row[1]), _s(row[2]), list(row[3:15])))
    return out


def _view_rows(ws, hdr: int, tc: int):
    for row in ws.iter_rows(min_row=hdr + 1, values_only=True):
        labels = tuple(_s(x) for x in row[:tc])
        if not any(labels) or any(x.strip() in TOTAL_LABELS for x in labels):
            continue
        yield labels, row


def view_labels(path: Path, kind: str) -> list[tuple[str, ...]]:
    from openpyxl import load_workbook
    sheet, hdr, tc, _ = VIEWS[kind]
    return [lab for lab, _ in _view_rows(load_workbook(path, read_only=True)[sheet], hdr, tc)]


def sumifs(rows, kind: str, labels: tuple[str, ...], day: date | None) -> list[float]:
    """The view's SUMIFS for one row, evaluated in Python with the formula's own criteria (exact text matches)."""
    crit = VIEWS[kind][3]
    tot = [0.0] * 12
    for d, div, off, vals in rows:
        if "date" in crit and d != day:
            continue
        if "division" in crit and div != labels[0]:
            continue
        if "office" in crit and off != labels[1]:
            continue
        for i, v in enumerate(vals[:12]):
            if isinstance(v, (int, float)):
                tot[i] += v
    return tot


def libreoffice_values(path: Path, kind: str, day: date | None, soffice: str, work: Path) -> dict[tuple[str, ...], list]:
    """Recalculate a copy in LibreOffice (B2 set to `day` for the daily sheets) and read the view's value rows."""
    from openpyxl import load_workbook
    sheet, hdr, tc, _ = VIEWS[kind]
    src = work / f"in_{kind}_{day or 'all'}.xlsx"
    wb = load_workbook(path)
    if day is not None:
        wb[sheet]["B2"] = datetime.combine(day, datetime.min.time())
    wb.save(src)  # openpyxl keeps no cached results, so LibreOffice has to calculate every formula
    out_dir = work / "out"
    out_dir.mkdir(exist_ok=True)
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out_dir), str(src)],
                   check=True, capture_output=True, timeout=600)
    ws = load_workbook(out_dir / src.name, read_only=True, data_only=True)[sheet]
    return {lab: [v if isinstance(v, (int, float)) else 0 for v in row[tc:tc + 12]] for lab, row in _view_rows(ws, hdr, tc)}


# --------------------------------------------------------------------------- our side
def our_tables(path: Path):
    from sqlalchemy import select

    from app.db import init_db, make_engine, make_session_factory
    from app.forms import analytics
    from app.forms.schema import create_form, recipient_codes, set_status
    from app.forms.service import save
    from app.models import Campaign
    from app.registry.loader import bootstrap
    from app.reporting.columns import COLUMN_KEYS, COLUMNS
    from scripts.import_sheet import read_sheet

    rows, problems = read_sheet(path)
    tmp = Path(tempfile.mkdtemp(prefix="parity_"))
    engine = make_engine(f"sqlite:///{(tmp / 'parity.sqlite3').as_posix()}")
    init_db(engine)
    s = make_session_factory(engine)()
    bootstrap(s, ROOT / "config")
    c = s.scalar(select(Campaign).where(Campaign.slug == "drought_measures"))
    form = create_form(s, "parity", {
        "title_mr": c.title_mr, "title_en": c.title_en, "frequency": "daily",
        "start_date": min(d for d, _, _ in rows).isoformat(),
        "fields": [{"key": x.key, "label_mr": x.label_mr, "label_en": x.label_en, "type": "number", "min_value": 0,
                    "group": x.group} for x in COLUMNS]})
    set_status(s, "parity", form, "published")
    for d, code, vals in rows:
        save(s, None, form, code, d, {k: v for k, v in zip(COLUMN_KEYS, vals) if v is not None}, confirm_zero=True,
             check=False, actor="import:xlsx", submit=True)
    s.flush()
    offices = recipient_codes(s, form)

    def table(kind: str, d0: date, d1: date) -> dict[tuple[str, ...], list]:
        t = analytics.table(s, form, offices, d0, d1, kind)
        tc = t["text_cols"]
        return {tuple(_s(x) for x in r[:tc]): r[tc:] for r in t["rows"]}

    def close():
        s.close()
        engine.dispose()
        shutil.rmtree(tmp, ignore_errors=True)
    return table, problems, close


# --------------------------------------------------------------------------- compare
def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--start", default="2026-10-09")
    ap.add_argument("--end", default="2026-10-15")
    ap.add_argument("--no-libreoffice", action="store_true")
    args = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    start, end = date.fromisoformat(args.start), date.fromisoformat(args.end)
    days = [start + timedelta(days=i) for i in range((end - start).days + 1)]

    from app.forms.sheet_names import OFFICE_MAP
    soffice = None if args.no_libreoffice else (shutil.which("soffice") or next(
        (p for p in (r"C:\Program Files\LibreOffice\program\soffice.exe", "/usr/bin/soffice",
                     "/Applications/LibreOffice.app/Contents/MacOS/soffice") if Path(p).exists()), None))
    print(f"workbook: {args.xlsx.name}\nrange: {start} … {end}\nLibreOffice: {soffice or 'not used'}")

    rows = data_rows(args.xlsx)
    whole = (min(d for d, *_ in rows), max(d for d, *_ in rows))
    table, problems, close = our_tables(args.xlsx)
    for p in problems:
        print("IMPORT PROBLEM:", p)

    work = Path(tempfile.mkdtemp(prefix="parity_lo_"))
    unexplained = explained = checked = lo_checked = lo_disagree = 0
    for kind, (sheet, _, tc, crit) in VIEWS.items():
        labels = view_labels(args.xlsx, kind)
        label_set = set(labels)
        # names in डेटा एंट्री that the view's SUMIFS can't match but that map to an office
        variant_codes = set()
        if kind != "datewise":
            for _, div, off, vals in rows:
                key = (div,) if tc == 1 else (div, off)
                if key not in label_set and any(isinstance(v, (int, float)) for v in vals) and (div, off) in OFFICE_MAP:
                    variant_codes.add(OFFICE_MAP[(div, off)])
        for d in (days if kind.startswith("daily_") else [None]):
            ours = table(kind, *((d, d) if d else whole))
            lo = libreoffice_values(args.xlsx, kind, d, soffice, work) if soffice else None
            for lab in labels:
                if kind == "datewise":
                    ld = _day(lab[0])
                    if ld is None or not start <= ld <= end:
                        continue
                    excel, key = sumifs(rows, kind, lab, ld), (ld.isoformat(),)
                else:
                    excel, key = sumifs(rows, kind, lab, d), lab
                if lo is not None:
                    lo_checked += 1
                    lv = lo.get(lab)
                    if lv is None or any(abs(x - y) > 1e-6 for x, y in zip(lv, excel)):
                        lo_disagree += 1
                        print(f"  LibreOffice ≠ SUMIFS: {sheet} {d or ''} {lab}: {lv} vs {excel}")
                mine = ours.get(key)
                if mine is None:
                    unexplained += 1
                    print(f"  MISSING ROW in ours: {sheet} {d or ''} {' / '.join(lab)}")
                    continue
                for i, (x, y) in enumerate(zip(excel, mine)):
                    checked += 1
                    if abs(x - (y or 0)) < 1e-6:
                        continue
                    why = ({OFFICE_MAP.get((lab[0], lab[1]))} & variant_codes) if tc == 2 else variant_codes
                    if why:
                        explained += 1
                    else:
                        unexplained += 1
                    tag = f"explained (name variant: {', '.join(sorted(why))})" if why else "UNEXPLAINED"
                    print(f"  {tag}: {sheet} {d or ''} {' / '.join(lab)} col {i + 1}: workbook {x:g}, ours {y}")
    print(f"\n{checked} cells compared · {explained} explained differences · {unexplained} unexplained"
          + (f" · LibreOffice agreed with SUMIFS on {lo_checked - lo_disagree}/{lo_checked} rows" if soffice else ""))
    close()
    shutil.rmtree(work, ignore_errors=True)
    sys.exit(1 if unexplained or lo_disagree else 0)


if __name__ == "__main__":
    main()
