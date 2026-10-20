"""The Excel workbook export: the admin's sheets in order, totals equal a hand-computed fixture, and admin
corrections are highlighted with the original value in a cell comment."""
from __future__ import annotations

import io
from datetime import date

from openpyxl import load_workbook

from app.forms import analytics
from app.forms.schema import create_form, set_status
from app.forms.service import correct, save
from app.models import User

D1, D2 = date(2026, 10, 1), date(2026, 10, 2)
ADMIN = User(username="admin_x", password_hash="", role="state")


def make(session):
    form = create_form(session, "t", {
        "title_mr": "दुष्काळ", "start_date": D1.isoformat(), "instructions_mr": "ओळ एक\nओळ दोन",
        "recipients": ["pune", "satara", "nashik"],
        "fields": [{"key": "articles", "label_mr": "लेख", "type": "number"},
                   {"key": "news", "label_mr": "वृत्त", "type": "number"},
                   {"key": "remark", "label_mr": "शेरा", "type": "text"}]})
    set_status(session, "t", form, "published")
    session.flush()
    # pune: 1/2 and 3/4; satara: 5/6 on day 1 only; nashik: day 2 draft (not counted)
    save(session, None, form, "pune", D1, {"articles": 1, "news": 2, "remark": "ठीक"}, check=False, submit=True)
    save(session, None, form, "pune", D2, {"articles": 3, "news": 4}, check=False, submit=True)
    save(session, None, form, "satara", D1, {"articles": 5, "news": 6}, check=False, submit=True)
    save(session, None, form, "nashik", D2, {"articles": 100}, check=False)
    correct(session, ADMIN, form, "satara", D1, "news", 60, "शून्य सुटले")
    session.flush()
    return form


def test_workbook_layout_totals_and_highlights(session):
    form = make(session)
    offices = ["pune", "satara", "nashik"]
    wb = load_workbook(io.BytesIO(analytics.workbook(session, form, offices, D1, D2)))
    assert wb.sheetnames == ["सूचना", "डेटा एंट्री", "दररोजचा अहवाल", "विभागनिहाय एकत्रित", "तारखेनिहाय अहवाल",
                             "जिल्हानिहाय दररोजचा अहवाल", "जिल्हानिहाय एकत्रित"]
    assert wb["सूचना"]["A3"].value == "ओळ एक" and wb["सूचना"]["A4"].value == "ओळ दोन"

    de = wb["डेटा एंट्री"]
    assert [c.value for c in de[3]] == ["दिनांक", "विभागाचे नाव", "जिल्ह्याचे नाव", "लेख", "वृत्त", "शेरा", "स्थिती"]
    rows = [[c.value for c in r] for r in de.iter_rows(min_row=4)]
    assert len(rows) == 2 * 3  # every office × day, blanks included
    by = {(r[0].date(), r[2]): r for r in rows}
    pune_name = rows[0][2]
    assert by[(D1, pune_name)][3:6] == [1, 2, "ठीक"]
    nashik = next(r for r in rows if r[0].date() == D2 and r[3:5] == [None, None] and r[6] == "मसुदा")
    assert nashik  # the draft row is shown blank, marked as a draft
    hl = [c for r in de.iter_rows(min_row=4) for c in r if c.fill.fill_type == "solid"]
    assert len(hl) == 1 and hl[0].value == 60 and hl[0].fill.fgColor.rgb.endswith("FFF2CC")
    assert "मूळ आकडा: 6" in hl[0].comment.text and "शून्य सुटले" in hl[0].comment.text

    cum = wb["जिल्हानिहाय एकत्रित"]
    total = [c.value for c in cum[cum.max_row]]
    assert total[1] == "एकत्रित एकूण" and total[2:4] == [1 + 3 + 5, 2 + 4 + 60]
    div = wb["विभागनिहाय एकत्रित"]
    assert [c.value for c in div[div.max_row]][:3] == ["राज्य एकूण", 9, 66]
    datewise = wb["तारखेनिहाय अहवाल"]
    assert [[c.value for c in r][:3] for r in datewise.iter_rows(min_row=4)] == \
        [[D1.isoformat(), 6, 62], [D2.isoformat(), 3, 4], ["एकूण", 9, 66]]
    daily = wb["दररोजचा अहवाल"]  # the last day of the range
    assert daily["A4"].value == "विभागाचे नाव" and [c.value for c in daily[daily.max_row]][:3] == ["राज्य एकूण", 3, 4]
    assert "२ ऑक्टोबर २०२६" in daily["A2"].value
