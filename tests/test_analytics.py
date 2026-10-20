"""Form analytics equal raw sums; scope; rankings and sheet tables; the real workbook import matches the sheet's
own totals."""
from __future__ import annotations

import io
import os
from datetime import date
from pathlib import Path

import pytest
from sqlalchemy import func, select

from app.auth import visible_office_codes
from app.forms import analytics
from app.forms.service import save
from app.models import FormValue, User
from app.reporting.aggregate import compute_report
from app.reporting.columns import COLUMN_KEYS

D1, D2 = date(2026, 10, 11), date(2026, 10, 12)
SHEET = Path(os.environ.get("DGIPR_SHEET", Path(__file__).resolve().parent.parent / "samples" /
                            "office_sheet_2026-10-15.xlsx"))


@pytest.fixture()
def filled(session, drought_form):
    vals = {"pune": 3, "satara": 5, "nashik": 7, "nashik_rio": 2}
    for office, base in vals.items():
        save(session, None, drought_form, office, D1, {k: base for k in COLUMN_KEYS}, check=False, actor="t",
             submit=True)
        save(session, None, drought_form, office, D2, {"articles": base * 10}, check=False, actor="t", submit=True)
    save(session, None, drought_form, "latur", D1, {"news": 1000}, check=False, actor="t")  # draft: never counted
    session.flush()
    return vals


def all_offices(session):
    return visible_office_codes(session, User(username="s", password_hash="", role="state"))


def test_totals_equal_raw_sums_of_submitted(session, drought_form, filled):
    offices = all_offices(session)
    t = analytics.totals(session, drought_form, offices, D1, D2)
    raw = session.scalar(select(func.sum(FormValue.value_num))) - 1000
    assert sum(t.values()) == raw
    assert t["articles"] == sum(filled.values()) * 11 and t["news"] == sum(filled.values())
    s = analytics.summary(session, drought_form, offices, D2, D2)
    assert s["previous"]["start"] == D1.isoformat() and s["previous"]["totals"]["news"] == sum(filled.values())
    assert s["coverage"]["offices_reporting"] == 4 and s["coverage"]["filled_cells"] == 4


def test_breakdown_division_and_office(session, drought_form, filled):
    offices = all_offices(session)
    div = analytics.breakdown(session, drought_form, offices, D1, D1, "division")
    by = {r["code"]: r["values"]["news"] for r in div["rows"]}
    assert by["pune"] == 8 and by["nashik"] == 9 and by["latur"] == 0 and div["total"]["news"] == 17
    off = analytics.breakdown(session, drought_form, offices, D1, D2, "office")
    assert len(off["rows"]) == 43
    assert next(r for r in off["rows"] if r["code"] == "satara")["values"]["articles"] == 55


def test_scope_limits_to_visible_offices(session, drought_form, filled):
    zone = User(username="z", password_hash="", role="zone", division_code="nashik")
    offices = visible_office_codes(session, zone)
    assert analytics.totals(session, drought_form, offices, D1, D1)["news"] == 9


def test_rankings_series_and_tables(session, drought_form, filled):
    offices = all_offices(session)
    r = analytics.rankings(session, drought_form, offices, D1, D1, "news")
    assert [x["code"] for x in r["rows"][:2]] == ["nashik", "satara"] and r["rows"][0]["share"] == round(700 / 17, 1)
    s = analytics.series(session, drought_form, offices, D1, D2)
    assert [x["values"]["articles"] for x in s] == [17, 170]
    t = analytics.table(session, drought_form, offices, D1, D2, "cumulative_office")
    assert t["total"][2:] == [analytics.totals(session, drought_form, offices, D1, D2)[k] for k in COLUMN_KEYS]
    assert "जिल्ह्याचे नाव" in analytics.to_csv(t) and t["text_cols"] == 2
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(analytics.to_xlsx(t))).active
    assert ws.cell(3, 2).value == "जिल्ह्याचे नाव" and ws.cell(ws.max_row, 3).value == 187
    daily = analytics.table(session, drought_form, offices, D2, D2, "daily_division")
    assert daily["total"][0] == "राज्य एकूण" and daily["start"] == daily["end"]


def test_export_endpoints(admin_client):
    r = admin_client.post("/api/forms", json={"title_mr": "निर्यात", "start_date": "2026-10-01",
                                              "fields": [{"label_mr": "लेख", "label_en": "Articles"}]})
    fid = r.json()["id"]
    r = admin_client.get(f"/api/forms/{fid}/analytics/export?kind=datewise&start=2026-10-01&end=2026-10-03")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    r = admin_client.get(f"/api/forms/{fid}/analytics/export?kind=daily_office&format=csv&start=2026-10-01")
    assert r.status_code == 200 and "text/csv" in r.headers["content-type"]
    r = admin_client.get(f"/api/forms/{fid}/export?start=2026-10-01&end=2026-10-03")
    assert r.status_code == 200 and r.content[:2] == b"PK" and "filename*=UTF-8" in r.headers["content-disposition"]


# --------------------------------------------------------------------------- the real workbook
@pytest.mark.skipif(not SHEET.is_file(), reason="office data-entry workbook not available")
def test_import_matches_sheet_totals(session, campaign, drought_form):
    from openpyxl import load_workbook

    from scripts.import_sheet import OFFICE_MAP, read_sheet
    rows, problems = read_sheet(SHEET)
    assert not problems
    for day, code, vals in rows:
        save(session, None, drought_form, code, day, {k: v for k, v in zip(COLUMN_KEYS, vals) if v is not None},
             confirm_zero=True, check=False, actor="import:xlsx", submit=True)
    session.flush()
    offices = all_offices(session)
    start, end = date(2026, 9, 26), date(2026, 10, 15)
    assert list(analytics.totals(session, drought_form, offices, start, end).values()) == \
           [8, 378, 4598, 57, 2, 790, 612, 738, 293, 80, 129, 0]
    # district cumulative totals equal the sheet's own "जिल्हानिहाय एकत्रित" figures
    wb = load_workbook(SHEET, read_only=True, data_only=True)
    expected = {OFFICE_MAP[(r[0], r[1])]: list(r[2:14])
                for r in wb["जिल्हानिहाय एकत्रित"].iter_rows(values_only=True) if (r[0], r[1]) in OFFICE_MAP}
    got = {r["code"]: [r["values"][k] for k in COLUMN_KEYS]
           for r in analytics.breakdown(session, drought_form, offices, start, end, "office")["rows"]}
    assert got == expected
    # the 15 Oct Social Monitoring report shows the sheet's state total (कात्रणे 820 that day)
    data = compute_report(session, campaign, date(2026, 10, 15), date(2026, 10, 15))
    assert data["grand_total"]["clippings"]["value"] == 820
